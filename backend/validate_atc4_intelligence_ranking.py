import csv
import json
import sqlite3
import sys
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]

DB_PATH = PROJECT_ROOT / "database" / "rxnorm_research.db"
INPUT_CSV = PROJECT_ROOT / "validation" / "atc4_ranking_expected.csv"
OUTPUT_XLSX = PROJECT_ROOT / "validation" / "atc4_ranking_validation_report.xlsx"
OUTPUT_JSON = PROJECT_ROOT / "validation" / "atc4_ranking_validation_results.json"


def safe_text(value: Any) -> str:
    return str(value or "").strip()


def clean_text(value: Any) -> str:
    return str(value or "").strip().lower()


def row_to_dict(row: sqlite3.Row) -> Dict[str, Any]:
    return {key: row[key] for key in row.keys()}


def fetch_classification_rows(conn: sqlite3.Connection, rxcui: str) -> List[Dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT *
        FROM research_classification_detail
        WHERE CAST(rxcui AS TEXT) = ?
          AND class_id IS NOT NULL
          AND TRIM(class_id) <> ''
        """,
        (str(rxcui),),
    ).fetchall()

    return [row_to_dict(row) for row in rows]


def build_primary_therapeutic_pathway_from_rows(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Sprint 15C.3 database validation copy of the ATC4 Intelligence Ranking Layer.
    Keep this logic aligned with build_primary_therapeutic_pathway() in rxnorm_intelligence_api_v1.py.
    """

    atc1 = [r for r in rows if r.get("class_type") == "ATC1"]
    atc2 = [r for r in rows if r.get("class_type") == "ATC2"]
    atc3 = [r for r in rows if r.get("class_type") == "ATC3"]
    atc4 = [r for r in rows if r.get("class_type") == "ATC4"]

    moa_rows = [r for r in rows if r.get("class_type") == "MOA"]
    epc_rows = [r for r in rows if r.get("class_type") == "EPC"]
    disease_rows = [r for r in rows if r.get("class_type") == "DISEASE"]
    va_rows = [r for r in rows if r.get("class_type") == "VA"]

    if not atc4:
        return {
            "primary_atc1": None,
            "primary_atc2": None,
            "primary_atc3": None,
            "primary_atc4": None,
            "pathway": [],
            "atc4_ranking": [],
            "ranking_version": "15C.3-db-validation",
        }

    def get_name(row: Dict[str, Any]) -> str:
        return safe_text(
            row.get("class_name")
            or row.get("atc_full_name")
            or row.get("class_id")
        )

    def clean_atc_name(row):
        if not row:
            return None

        row = dict(row)
        class_id = safe_text(row.get("class_id"))

        atc_name_fallbacks = {
            "A": "Alimentary Tract and Metabolism",
            "A10": "Drugs Used in Diabetes",
            "A10B": "Blood Glucose Lowering Drugs, Excluding Insulins",
            "A10BJ": "Glucagon-like Peptide-1 (GLP-1) Analogs",
            "C": "Cardiovascular System",
            "M": "Musculo-Skeletal System",
            "M01": "Anti-inflammatory and antirheumatic products",
            "M01A": "Anti-inflammatory and antirheumatic products, non-steroids",
            "M01AE": "Propionic acid derivatives",
            "N": "Nervous System",
            "R": "Respiratory System",
            "G": "Genito Urinary System and Sex Hormones",
            "J": "Antiinfectives for Systemic Use",
        }

        row["class_name"] = (
            row.get("class_name")
            or atc_name_fallbacks.get(class_id)
            or row.get("atc_full_name")
            or class_id
            or ""
        )

        return row

    evidence_text = " ".join(
        [
            " ".join(get_name(r) for r in moa_rows),
            " ".join(get_name(r) for r in epc_rows),
            " ".join(get_name(r) for r in disease_rows),
            " ".join(get_name(r) for r in va_rows),
        ]
    ).lower()

    def score_atc4_candidate(row: Dict[str, Any]) -> Dict[str, Any]:
        code = safe_text(row.get("class_id"))
        name = clean_text(get_name(row))
        full_name = clean_text(row.get("atc_full_name"))
        combined = f"{code.lower()} {name} {full_name}"

        score = 0
        reasons = []

        score += len(code) * 3
        reasons.append(f"specificity_bonus={len(code) * 3}")

        if row.get("is_atc_hierarchy_row"):
            score += 20
            reasons.append("official_atc_hierarchy_row=20")

        nsaid_evidence = any(
            term in evidence_text
            for term in [
                "cyclooxygenase",
                "nonsteroidal",
                "anti-inflammatory",
                "antiinflammatory",
                "pain",
                "inflammation",
                "arthritis",
                "bursitis",
                "fever",
                "osteoarthritis",
                "rheumatoid",
            ]
        )

        if nsaid_evidence:
            if code.startswith("M01AE"):
                score += 160
                reasons.append("nsaid_propionic_acid_alignment=160")
            elif code.startswith("M01A"):
                score += 110
                reasons.append("nsaid_antiinflammatory_alignment=110")
            elif code.startswith("M02A"):
                score += 60
                reasons.append("topical_antiinflammatory_secondary=60")
            elif code.startswith("N02"):
                score += 35
                reasons.append("analgesic_secondary_pathway=35")

        if "propionic acid" in combined:
            score += 90
            reasons.append("propionic_acid_name_match=90")

        metabolic_evidence = any(
            term in evidence_text
            for term in [
                "diabetes",
                "glucose",
                "hyperglycemia",
                "metabolic",
                "glucagon-like",
                "glp",
                "incretin",
                "weight",
                "obesity",
            ]
        )
        if "proton pump" in evidence_text or "gastric acid" in evidence_text or "gastroesophageal reflux" in evidence_text:
            if code.startswith("A02BC"):
                score += 250
                reasons.append("ppi_primary_pathway_bonus=250")
            elif code.startswith("A02BD"):
                score -= 100
                reasons.append("h_pylori_combination_penalty=-100")
            elif code.startswith("B01AC"):
                score -= 150
                reasons.append("non_primary_antiplatelet_penalty_for_ppi=-150")
                
        if metabolic_evidence:
            if code.startswith("A10BJ"):
                score += 180
                reasons.append("glp1_diabetes_alignment=180")
            elif code.startswith("A10B"):
                score += 120
                reasons.append("blood_glucose_lowering_alignment=120")
            elif code.startswith("A10"):
                score += 90
                reasons.append("diabetes_drug_alignment=90")
            elif code.startswith("A"):
                score += 50
                reasons.append("metabolism_domain_alignment=50")

        if "glucagon-like peptide" in combined or "glp-1" in combined or "glp1" in combined:
            score += 120
            reasons.append("glp1_name_match=120")

        if "blood glucose" in combined or "diabetes" in combined:
            score += 80
            reasons.append("diabetes_name_match=80")

        cardiovascular_evidence = any(
            term in evidence_text
            for term in [
                "cardiovascular",
                "platelet",
                "thrombosis",
                "myocardial",
                "stroke",
                "infarction",
                "ischemic",
                "antiplatelet",
            ]
        )

        if cardiovascular_evidence:
            if code.startswith("C"):
                score += 140
                reasons.append("cardiovascular_alignment=140")
            if "antithrombotic" in combined or "platelet" in combined:
                score += 100
                reasons.append("antithrombotic_name_match=100")

        combination_terms = [
            "combination",
            "combinations",
            "in combination",
            "with other drugs",
        ]

        if any(term in combined for term in combination_terms):
            score -= 200
            reasons.append("combination_product_penalty=-200")

        if code.startswith("J01"):
            score += 40
            reasons.append("systemic_antiinfective_bonus=40")

        drug_name_text = clean_text(row.get("related_name")) + " " + clean_text(row.get("drug_name"))

        if code.startswith("B01AC") and "aspirin" in drug_name_text:
            score += 50
            reasons.append("aspirin_antiplatelet_pathway_bonus=50")

        if "biguanide" in combined:
            score += 250

        if nsaid_evidence and code.startswith(("R", "G", "C01")):
            score -= 35
            reasons.append("secondary_route_penalty=-35")

        if metabolic_evidence and not code.startswith("A"):
            score -= 50
            reasons.append("non_metabolic_penalty=-50")

        return {
            "class_id": code,
            "class_name": get_name(row),
            "score": score,
            "reasons": reasons,
        }

    ranked_atc4 = sorted(
        atc4,
        key=lambda row: (
            score_atc4_candidate(row)["score"],
            len(safe_text(row.get("class_id"))),
            safe_text(row.get("class_id")),
        ),
        reverse=True,
    )

    primary_atc4 = ranked_atc4[0]
    atc4_code = safe_text(primary_atc4.get("class_id"))

    atc3_code = atc4_code[:4]
    atc2_code = atc4_code[:3]
    atc1_code = atc4_code[:1]

    primary_atc3 = next(
        (row for row in atc3 if safe_text(row.get("class_id")) == atc3_code),
        None,
    )
    primary_atc2 = next(
        (row for row in atc2 if safe_text(row.get("class_id")) == atc2_code),
        None,
    )
    primary_atc1 = next(
        (row for row in atc1 if safe_text(row.get("class_id")) == atc1_code),
        None,
    )

    primary_atc1 = clean_atc_name(primary_atc1)
    primary_atc2 = clean_atc_name(primary_atc2)
    primary_atc3 = clean_atc_name(primary_atc3)
    primary_atc4 = clean_atc_name(primary_atc4)

    pathway = [
        item
        for item in [primary_atc1, primary_atc2, primary_atc3, primary_atc4]
        if item
    ]

    atc4_ranking = [score_atc4_candidate(row) for row in ranked_atc4[:10]]

    return {
        "primary_atc1": primary_atc1,
        "primary_atc2": primary_atc2,
        "primary_atc3": primary_atc3,
        "primary_atc4": primary_atc4,
        "pathway": pathway,
        "atc4_ranking": atc4_ranking,
        "ranking_version": "15C.3-db-validation",
    }


def read_expected_csv() -> List[Dict[str, str]]:
    if not INPUT_CSV.exists():
        raise FileNotFoundError(f"Expected input file not found: {INPUT_CSV}")

    with INPUT_CSV.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)

        required_columns = {
            "rxcui",
            "drug_name",
            "expected_atc1",
            "expected_atc2",
            "expected_atc3",
            "expected_atc4",
        }

        missing = required_columns - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Input CSV is missing required columns: {sorted(missing)}")

        return [dict(row) for row in reader]


def compare_expected_actual(expected: Dict[str, str], pathway: Dict[str, Any]) -> Dict[str, Any]:
    primary_atc1 = pathway.get("primary_atc1") or {}
    primary_atc2 = pathway.get("primary_atc2") or {}
    primary_atc3 = pathway.get("primary_atc3") or {}
    primary_atc4 = pathway.get("primary_atc4") or {}

    actual = {
        "actual_atc1": safe_text(primary_atc1.get("class_id")),
        "actual_atc2": safe_text(primary_atc2.get("class_id")),
        "actual_atc3": safe_text(primary_atc3.get("class_id")),
        "actual_atc4": safe_text(primary_atc4.get("class_id")),
        "actual_atc1_name": safe_text(primary_atc1.get("class_name")),
        "actual_atc2_name": safe_text(primary_atc2.get("class_name")),
        "actual_atc3_name": safe_text(primary_atc3.get("class_name")),
        "actual_atc4_name": safe_text(primary_atc4.get("class_name")),
        "ranking_version": safe_text(pathway.get("ranking_version")),
        "atc4_ranking_json": json.dumps(pathway.get("atc4_ranking") or [], ensure_ascii=False),
    }

    checks = {
        "atc1_match": safe_text(expected.get("expected_atc1")) == actual["actual_atc1"],
        "atc2_match": safe_text(expected.get("expected_atc2")) == actual["actual_atc2"],
        "atc3_match": safe_text(expected.get("expected_atc3")) == actual["actual_atc3"],
        "atc4_match": safe_text(expected.get("expected_atc4")) == actual["actual_atc4"],
    }

    return {
        **actual,
        **checks,
        "match_status": "PASS" if all(checks.values()) else "FAIL",
    }


def run_validation() -> List[Dict[str, Any]]:
    expected_rows = read_expected_csv()

    if not DB_PATH.exists():
        raise FileNotFoundError(f"SQLite database not found: {DB_PATH}")

    results: List[Dict[str, Any]] = []

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    try:
        for index, expected in enumerate(expected_rows, start=1):
            rxcui = safe_text(expected.get("rxcui"))
            drug_name = safe_text(expected.get("drug_name"))

            print(f"[{index}/{len(expected_rows)}] Validating RxCUI {rxcui} - {drug_name}")

            classification_rows = fetch_classification_rows(conn, rxcui)

            if not classification_rows:
                results.append(
                    {
                        "rxcui": rxcui,
                        "drug_name": drug_name,
                        "expected_atc1": expected.get("expected_atc1"),
                        "expected_atc2": expected.get("expected_atc2"),
                        "expected_atc3": expected.get("expected_atc3"),
                        "expected_atc4": expected.get("expected_atc4"),
                        "actual_atc1": "",
                        "actual_atc2": "",
                        "actual_atc3": "",
                        "actual_atc4": "",
                        "actual_atc1_name": "",
                        "actual_atc2_name": "",
                        "actual_atc3_name": "",
                        "actual_atc4_name": "",
                        "atc1_match": False,
                        "atc2_match": False,
                        "atc3_match": False,
                        "atc4_match": False,
                        "match_status": "ERROR",
                        "ranking_version": "",
                        "atc4_ranking_json": "",
                        "error_message": "No classification rows found for RxCUI",
                    }
                )
                continue

            pathway = build_primary_therapeutic_pathway_from_rows(classification_rows)
            comparison = compare_expected_actual(expected, pathway)

            results.append(
                {
                    "rxcui": rxcui,
                    "drug_name": drug_name,
                    "expected_atc1": expected.get("expected_atc1"),
                    "expected_atc2": expected.get("expected_atc2"),
                    "expected_atc3": expected.get("expected_atc3"),
                    "expected_atc4": expected.get("expected_atc4"),
                    **comparison,
                    "error_message": "",
                }
            )

    finally:
        conn.close()

    return results


def export_json(results: List[Dict[str, Any]]) -> None:
    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT_JSON.open("w", encoding="utf-8") as file:
        json.dump(results, file, indent=2, ensure_ascii=False)


def export_excel(results: List[Dict[str, Any]]) -> None:
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "ATC4 Ranking QA"

    headers = [
        "rxcui",
        "drug_name",
        "expected_atc1",
        "actual_atc1",
        "atc1_match",
        "expected_atc2",
        "actual_atc2",
        "atc2_match",
        "expected_atc3",
        "actual_atc3",
        "atc3_match",
        "expected_atc4",
        "actual_atc4",
        "atc4_match",
        "match_status",
        "actual_atc1_name",
        "actual_atc2_name",
        "actual_atc3_name",
        "actual_atc4_name",
        "ranking_version",
        "error_message",
        "atc4_ranking_json",
    ]

    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for result in results:
        ws.append([result.get(header, "") for header in headers])

    match_status_col = headers.index("match_status") + 1

    for row in ws.iter_rows(min_row=2):
        status = row[match_status_col - 1].value

        if status == "PASS":
            status_fill = pass_fill
        elif status == "FAIL":
            status_fill = fail_fill
        else:
            status_fill = warn_fill

        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if cell.column == match_status_col:
                cell.fill = status_fill

    for col_idx, header in enumerate(headers, start=1):
        max_length = len(header)

        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            max_length = max(max_length, len(str(value or "")))

        if header == "atc4_ranking_json":
            width = 80
        else:
            width = min(max_length + 2, 45)

        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    total = len(results)
    passed = sum(1 for item in results if item.get("match_status") == "PASS")
    failed = sum(1 for item in results if item.get("match_status") == "FAIL")
    errors = sum(1 for item in results if item.get("match_status") == "ERROR")
    pass_rate = round((passed / total) * 100, 1) if total else 0

    summary_rows = [
        ["Metric", "Value"],
        ["Database", str(DB_PATH)],
        ["Input CSV", str(INPUT_CSV)],
        ["Output XLSX", str(OUTPUT_XLSX)],
        ["Total Tested", total],
        ["Passed", passed],
        ["Failed", failed],
        ["Errors", errors],
        ["Pass Rate", f"{pass_rate}%"],
    ]

    for row in summary_rows:
        summary.append(row)

    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font

    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 100

    wb.save(OUTPUT_XLSX)


def main() -> None:
    print("==============================================")
    print("Sprint 15C.3 ATC4 Intelligence Ranking QA")
    print("Database-driven validation")
    print("==============================================")
    print(f"Database:    {DB_PATH}")
    print(f"Input CSV:   {INPUT_CSV}")
    print(f"Output XLSX: {OUTPUT_XLSX}")
    print("")

    try:
        results = run_validation()
        export_json(results)
        export_excel(results)

        total = len(results)
        passed = sum(1 for item in results if item.get("match_status") == "PASS")
        failed = sum(1 for item in results if item.get("match_status") == "FAIL")
        errors = sum(1 for item in results if item.get("match_status") == "ERROR")

        print("")
        print("==============================================")
        print("Validation Complete")
        print("==============================================")
        print(f"Total:  {total}")
        print(f"Passed: {passed}")
        print(f"Failed: {failed}")
        print(f"Errors: {errors}")
        print(f"Excel:  {OUTPUT_XLSX}")
        print(f"JSON:   {OUTPUT_JSON}")

        if failed or errors:
            sys.exit(1)

    except Exception as error:
        print("")
        print("VALIDATION FAILED TO RUN")
        print(f"{type(error).__name__}: {error}")
        sys.exit(1)


if __name__ == "__main__":
    main()